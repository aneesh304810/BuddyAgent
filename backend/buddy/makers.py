"""Buddy makes things you can see and keep: diagrams, charts, images, Word documents and
PowerPoint decks. They show up in the chat as cards (a picture, or a file to tap and open).

  make_diagram    a flowchart, architecture, mind map or org chart (Graphviz)
  make_chart      bar, line or pie charts from numbers (matplotlib)
  make_document   a Word document or PowerPoint deck from an outline (python-docx / python-pptx);
                  diagrams, charts and images made earlier can go inside
  make_image      an AI image (Gemini's image model, with your Gemini key); general prompts only,
                  like the other free models: personal details are refused

Nothing here runs code from the model: Claude gives structured content (or Graphviz text,
checked for anything that reads files) and Buddy builds the file. Files are kept in Buddy's data
folder; links to them are signed and expire, so only you can open them.
"""

import base64
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import subprocess
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path

from . import cards
from .config import settings
from .memory import Memory
from .tools import ToolError

FILES_DIR = Path(os.getenv("BUDDY_FILES_DIR", str(Path(settings.db_path).parent / "files")
                           if settings.db_path != ":memory:" else str(Path(tempfile.gettempdir()) / "buddy-files")))
LINK_DAYS = 30
MAX_FILE_BYTES = 25_000_000
MIME = {"png": "image/png", "jpg": "image/jpeg", "pdf": "application/pdf", "csv": "text/csv",
        "txt": "text/plain", "md": "text/markdown", "json": "application/json", "zip": "application/zip",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}
SCHEMA = """CREATE TABLE IF NOT EXISTS files (
    id TEXT PRIMARY KEY, name TEXT NOT NULL, ext TEXT NOT NULL, size INTEGER NOT NULL,
    kind TEXT NOT NULL, created_at TEXT NOT NULL)"""


# --- storage and signed links ------------------------------------------------------
def _secret() -> bytes:
    return (os.environ.get("BUDDY_TOKEN") or "dev-only").encode()


def sign(file_id: str, expires: int) -> str:
    return hmac.new(_secret(), f"{file_id}:{expires}".encode(), hashlib.sha256).hexdigest()[:32]


def link(file_id: str) -> str:
    expires = int(time.time()) + LINK_DAYS * 86400
    return f"/api/files/{file_id}?exp={expires}&sig={sign(file_id, expires)}"


def link_ok(file_id: str, exp: str, sig: str) -> bool:
    try:
        expires = int(exp)
    except (TypeError, ValueError):
        return False
    return expires > time.time() and hmac.compare_digest(sign(file_id, expires), sig or "")


def save(memory: Memory, data: bytes, name: str, ext: str, kind: str) -> dict:
    if len(data) > MAX_FILE_BYTES:
        raise ToolError("That file came out too big.")
    memory.db.execute(SCHEMA)
    FILES_DIR.mkdir(parents=True, exist_ok=True, mode=0o700)
    file_id = secrets.token_hex(12)
    (FILES_DIR / f"{file_id}.{ext}").write_bytes(data)
    clean = re.sub(r"[^\w .()-]+", "", name).strip()[:80] or "Buddy"
    memory.db.execute("INSERT INTO files (id, name, ext, size, kind, created_at) VALUES (?, ?, ?, ?, ?, ?)",
                      (file_id, clean, ext, len(data), kind, datetime.now(timezone.utc).isoformat()))
    memory.db.commit()
    return {"id": file_id, "name": f"{clean}.{ext}", "ext": ext, "size": len(data), "kind": kind}


def find(memory: Memory, file_id: str) -> tuple[Path, dict] | None:
    if not re.fullmatch(r"[0-9a-f]{24}", file_id or ""):
        return None
    memory.db.execute(SCHEMA)
    row = memory.db.execute("SELECT * FROM files WHERE id = ?", (file_id,)).fetchone()
    if not row:
        return None
    path = FILES_DIR / f"{file_id}.{row['ext']}"
    return (path, dict(row)) if path.exists() else None


def _card(file: dict, title: str, note: str = "") -> dict:
    url = link(file["id"])
    if file["ext"] in ("png", "jpg"):
        return {"kind": "image", "title": title, "subtitle": file["kind"], "image": url, "url": url, "note": note,
                "file_id": file["id"]}
    size = f"{file['size'] / 1024:.0f} KB" if file["size"] < 1_000_000 else f"{file['size'] / 1e6:.1f} MB"
    label = {"docx": "Word document", "pptx": "PowerPoint deck", "xlsx": "Excel spreadsheet", "pdf": "PDF",
             "csv": "CSV data", "txt": "Text", "md": "Text", "json": "JSON data", "zip": "Zip archive"
             }.get(file["ext"], file["ext"])
    return {"kind": "file", "title": file["name"], "subtitle": f"{label} · {size}", "url": url, "note": note,
            "file_id": file["id"], "ext": file["ext"]}


def card_for(memory: Memory, file_id: str, note: str = "") -> dict | None:
    """The chat card for a file Buddy keeps (for results delivered later, like a finished task)."""
    found = find(memory, file_id)
    if not found:
        return None
    row = found[1]
    file = {"id": row["id"], "name": f"{row['name']}.{row['ext']}", "ext": row["ext"], "size": row["size"],
            "kind": row["kind"]}
    return _card(file, row["name"] if row["ext"] in ("png", "jpg") else file["name"], note)


def _show(card: dict) -> str:
    """Made things appear in the chat by themselves (no show_cards step needed)."""
    box = cards.current.get()
    if box is not None:
        card_id = box.add(card)
        box.picked.append({**box.cards[card_id], "note": card.get("note", "")})
        return card_id
    return ""


# --- diagrams ----------------------------------------------------------------------
UNSAFE_DOT = re.compile(r"(?i)\b(image|imagepath|shapefile|fontpath|fontname\s*=\s*\"?/|stylesheet|href|url|URL)\s*=")
STYLE = ('graph [bgcolor="white" fontname="Helvetica" pad="0.3" nodesep="0.5" ranksep="0.6"];\n'
         'node [shape=box style="rounded,filled" fillcolor="#EEF3FF" color="#5B8CFF" fontname="Helvetica" '
         'fontsize=12 margin="0.2,0.1"];\nedge [color="#6B7280" fontname="Helvetica" fontsize=10];\n')


def render_diagram(dot: str) -> bytes:
    dot = dot.strip()
    if not re.match(r"(?is)^(strict\s+)?(di)?graph\b", dot) or len(dot) > 40_000:
        raise ToolError("Give Graphviz text starting with 'digraph' or 'graph'.")
    if UNSAFE_DOT.search(dot):
        raise ToolError("Diagrams can't use images, links or file paths.")
    styled = re.sub(r"\{", "{\n" + STYLE, dot, count=1)
    try:
        done = subprocess.run(["dot", "-Tpng", "-Gdpi=160"], input=styled.encode(), capture_output=True, timeout=20)
    except FileNotFoundError as e:
        raise ToolError("Graphviz isn't installed on the server (sudo apt install graphviz).") from e
    except subprocess.TimeoutExpired as e:
        raise ToolError("That diagram took too long to draw: make it simpler.") from e
    if done.returncode != 0:
        raise ToolError("The diagram has a mistake: " + done.stderr.decode(errors="replace")[-300:])
    return done.stdout


# --- charts ------------------------------------------------------------------------
PALETTE = ["#4C78DB", "#F2A93B", "#3DB389", "#D55E5E", "#8E6BD8", "#5BB8D9", "#B07A4A", "#7D8597"]


def render_chart(spec: dict) -> bytes:
    os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "buddy-mpl"))
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    labels = [str(x) for x in spec.get("labels", [])][:60]
    series = spec.get("series") or []
    if not labels or not series:
        raise ToolError("A chart needs labels and at least one series of numbers.")
    kind = spec.get("type", "bar")
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=150)
    try:
        if kind == "pie":
            values = [float(v) for v in series[0]["values"][:len(labels)]]
            ax.pie(values, labels=labels, colors=PALETTE[:len(values)], autopct="%1.0f%%", startangle=90,
                   wedgeprops={"edgecolor": "white", "linewidth": 2})
            ax.axis("equal")
        else:
            width = 0.8 / max(1, len(series))
            for i, s in enumerate(series[:8]):
                values = [float(v) for v in s["values"][:len(labels)]]
                if kind == "line":
                    ax.plot(labels[:len(values)], values, marker="o", linewidth=2, color=PALETTE[i], label=s.get("name"))
                else:
                    xs = [x + i * width - 0.4 + width / 2 for x in range(len(values))]
                    ax.bar(xs, values, width=width * 0.92, color=PALETTE[i], label=s.get("name"))
            if kind != "line":
                ax.set_xticks(range(len(labels)))
                ax.set_xticklabels(labels, rotation=30 if len(labels) > 6 else 0, ha="right" if len(labels) > 6 else "center")
            ax.spines[["top", "right"]].set_visible(False)
            ax.grid(axis="y", alpha=0.25)
            if spec.get("unit"):
                ax.set_ylabel(spec["unit"])
            if len(series) > 1:
                ax.legend(frameon=False)
        ax.set_title(spec.get("title", ""), fontsize=13, fontweight="bold", loc="left")
        fig.tight_layout()
        out = io.BytesIO()
        fig.savefig(out, format="png", facecolor="white")
        return out.getvalue()
    except (KeyError, ValueError, TypeError) as e:
        raise ToolError(f"The chart numbers don't fit together ({e}).") from e
    finally:
        plt.close(fig)


# --- documents and decks -------------------------------------------------------------
def _image_bytes(memory: Memory, file_id: str) -> bytes | None:
    found = find(memory, file_id or "")
    return found[0].read_bytes() if found and found[1]["ext"] in ("png", "jpg") else None


def build_docx(memory: Memory, spec: dict) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(11)
    doc.add_heading(spec.get("title", "Document"), level=0)
    if spec.get("subtitle"):
        p = doc.add_paragraph(spec["subtitle"])
        p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    for section in spec.get("sections", [])[:60]:
        if section.get("heading"):
            doc.add_heading(section["heading"], level=1)
        for para in section.get("paragraphs", []):
            doc.add_paragraph(str(para))
        for bullet in section.get("bullets", []):
            doc.add_paragraph(str(bullet), style="List Bullet")
        table = section.get("table")
        if table and table.get("columns"):
            t = doc.add_table(rows=1, cols=len(table["columns"]))
            t.style = "Light Grid Accent 1"
            for cell, name in zip(t.rows[0].cells, table["columns"]):
                cell.text = str(name)
            for row in table.get("rows", [])[:200]:
                cells = t.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(value)
        image = _image_bytes(memory, section.get("image_file_id", ""))
        if image:
            doc.add_picture(io.BytesIO(image), width=Inches(6))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def build_pptx(memory: Memory, spec: dict) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # widescreen
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = spec.get("title", "Presentation")
    if spec.get("subtitle"):
        title_slide.placeholders[1].text = spec["subtitle"]
    for section in spec.get("sections", [])[:40]:
        image = _image_bytes(memory, section.get("image_file_id", ""))
        table = section.get("table")
        slide = prs.slides.add_slide(prs.slide_layouts[5 if (image or table) else 1])
        slide.shapes.title.text = section.get("heading", "")
        slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)
        points = [str(p) for p in section.get("bullets", [])] + [str(p) for p in section.get("paragraphs", [])]
        if image or table:
            left = Inches(0.6)
            if points:
                box = slide.shapes.add_textbox(left, Inches(1.6), Inches(5.2), Inches(5)).text_frame
                box.word_wrap = True
                for i, point in enumerate(points[:8]):
                    para = box.paragraphs[0] if i == 0 else box.add_paragraph()
                    para.text, para.font.size = "• " + point, Pt(18)
                left = Inches(6.1)
            width = Inches(12.7) - left
            if image:
                slide.shapes.add_picture(io.BytesIO(image), left, Inches(1.6), width=width)
            elif table and table.get("columns"):
                rows = table.get("rows", [])[:12]
                shape = slide.shapes.add_table(len(rows) + 1, len(table["columns"]), left, Inches(1.6), width,
                                               Inches(0.4) * (len(rows) + 1)).table
                for c, name in enumerate(table["columns"]):
                    shape.cell(0, c).text = str(name)
                for r, row in enumerate(rows, start=1):
                    for c, value in enumerate(row[:len(table["columns"])]):
                        shape.cell(r, c).text = str(value)
        else:
            body = slide.placeholders[1].text_frame
            for i, point in enumerate(points[:10]):
                para = body.paragraphs[0] if i == 0 else body.add_paragraph()
                para.text, para.font.size = point, Pt(22)
        if section.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(section["notes"])
    for slide in prs.slides:  # a consistent accent colour on titles
        if slide.shapes.title is not None:
            for para in slide.shapes.title.text_frame.paragraphs:
                para.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# --- AI images ---------------------------------------------------------------------
IMAGE_MODEL = os.getenv("BUDDY_IMAGE_MODEL", "gemini-2.5-flash-image")


def generate_image(prompt: str, http=None) -> bytes:
    """An AI image from Gemini's image model, with your Gemini key (header, not in the URL)."""
    import urllib.request

    key = (os.environ.get("BUDDY_GEMINI_KEY") or "").strip()
    if not key:
        raise ToolError("Images need a Gemini key (aistudio.google.com → Get API key; set.sh BUDDY_GEMINI_KEY).")
    body = json.dumps({"contents": [{"parts": [{"text": prompt}]}],
                       "generationConfig": {"responseModalities": ["TEXT", "IMAGE"]}}).encode()
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{IMAGE_MODEL}:generateContent"
    request = urllib.request.Request(url, data=body, method="POST",
                                     headers={"Content-Type": "application/json", "x-goog-api-key": key})
    try:
        if http:
            data = http(request)
        else:
            with urllib.request.urlopen(request, timeout=90) as r:  # https only  # nosec B310
                data = json.loads(r.read())
    except Exception as e:
        raise ToolError(f"The image model didn't answer ({type(e).__name__}); BUDDY_IMAGE_MODEL may need "
                        "updating, or the free limit is reached.") from e
    for candidate in data.get("candidates", []):
        for part in candidate.get("content", {}).get("parts", []):
            inline = part.get("inlineData") or part.get("inline_data")
            if inline and inline.get("data"):
                return base64.b64decode(inline["data"])
    raise ToolError("The image model didn't return an image (it may have declined the prompt).")


# --- tools -------------------------------------------------------------------------
SECTION = {
    "type": "object",
    "properties": {
        "heading": {"type": "string"},
        "paragraphs": {"type": "array", "items": {"type": "string"}},
        "bullets": {"type": "array", "items": {"type": "string"}},
        "table": {"type": "object", "properties": {
            "columns": {"type": "array", "items": {"type": "string"}},
            "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}}}},
        "image_file_id": {"type": "string", "description": "A diagram, chart or image you made earlier"},
        "notes": {"type": "string", "description": "Speaker notes (slides)"},
    },
}
MAKER_TOOLS = [
    {"name": "make_diagram",
     "description": "Draw a diagram (flowchart, process, architecture, mind map, org chart, timeline) as a "
     "picture in the chat. Give Graphviz DOT text: digraph { rankdir=LR; a -> b; } with clear, short labels. "
     "Styling is added for you.",
     "input_schema": {"type": "object", "properties": {"title": {"type": "string"}, "dot": {"type": "string"}},
                      "required": ["title", "dot"], "additionalProperties": False}},
    {"name": "make_chart",
     "description": "Draw a bar, line or pie chart as a picture in the chat, from real numbers you have.",
     "input_schema": {"type": "object", "properties": {
         "title": {"type": "string"}, "type": {"type": "string", "enum": ["bar", "line", "pie"]},
         "labels": {"type": "array", "items": {"type": "string"}},
         "series": {"type": "array", "items": {"type": "object", "properties": {
             "name": {"type": "string"}, "values": {"type": "array", "items": {"type": "number"}}},
             "required": ["values"]}},
         "unit": {"type": "string"}},
         "required": ["title", "type", "labels", "series"], "additionalProperties": False}},
    {"name": "make_document",
     "description": "Write a Word document (docx) or PowerPoint deck (pptx) the user can open, share and edit. "
     "Give the whole content as an outline: sections with a heading, paragraphs, bullets, an optional table, "
     "and optionally a diagram, chart or image you made earlier (image_file_id). For slides, one section per "
     "slide, 3-6 short bullets each, speaker notes if useful. Appears in the chat as a file to open.",
     "input_schema": {"type": "object", "properties": {
         "format": {"type": "string", "enum": ["docx", "pptx"]}, "title": {"type": "string"},
         "subtitle": {"type": "string"}, "sections": {"type": "array", "items": SECTION}},
         "required": ["format", "title", "sections"], "additionalProperties": False}},
    {"name": "make_image",
     "description": "Create an AI image (illustration, picture, logo idea, scene) from a detailed prompt; it "
     "appears in the chat. General prompts only: nothing personal about the user (it's refused).",
     "input_schema": {"type": "object", "properties": {"prompt": {"type": "string"}, "title": {"type": "string"}},
                      "required": ["prompt"], "additionalProperties": False}},
]
MAKER_TOOL_NAMES = {t["name"] for t in MAKER_TOOLS}


async def run_maker_tool(memory: Memory, name: str, args: dict) -> str:
    import asyncio

    title = str(args.get("title") or "Buddy")
    if name == "make_diagram":
        data = await asyncio.to_thread(render_diagram, str(args.get("dot", "")))
        file = save(memory, data, title, "png", "diagram")
    elif name == "make_chart":
        data = await asyncio.to_thread(render_chart, args)
        file = save(memory, data, title, "png", "chart")
    elif name == "make_document":
        fmt = args.get("format")
        if fmt not in ("docx", "pptx"):
            raise ToolError("format is docx or pptx.")
        build = build_docx if fmt == "docx" else build_pptx
        data = await asyncio.to_thread(build, memory, args)
        file = save(memory, data, title, fmt, "document" if fmt == "docx" else "slides")
    elif name == "make_image":
        from .free_models import check_general

        prompt = str(args.get("prompt", ""))[:4000]
        check_general(prompt, memory)  # raises PersonalData: personal details never leave Claude
        data = await asyncio.to_thread(generate_image, prompt)
        file = save(memory, data, title or prompt[:40], "png", "image")
    else:
        raise ToolError(f"Unknown tool {name}")
    card_id = _show(_card(file, title))
    return json.dumps({"made": file["name"], "file_id": file["id"], "card": card_id,
                       "note": "It's in the chat already; just say a few words about it."})
