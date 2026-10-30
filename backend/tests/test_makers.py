"""Diagrams, charts, images, Word documents and decks, shown in the chat."""

import asyncio
import base64
import io
import json
import shutil
import subprocess
import time

import pytest

from buddy import cards, makers
from buddy.free_models import PersonalData
from buddy.memory import Memory
from buddy.tools import ToolError

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


@pytest.fixture
def mem(tmp_path, monkeypatch):
    monkeypatch.setattr(makers, "FILES_DIR", tmp_path / "files")
    monkeypatch.setenv("BUDDY_TOKEN", "t" * 48)
    return Memory(":memory:")


def run(mem, name, args):
    box = cards.CardBox()
    token = cards.current.set(box)
    try:
        return json.loads(asyncio.run(makers.run_maker_tool(mem, name, args))), box
    finally:
        cards.current.reset(token)


def test_chart_appears_in_the_chat_with_a_signed_link(mem):
    out, box = run(mem, "make_chart", {"title": "Spending", "type": "bar", "labels": ["Food", "Travel"],
                                       "series": [{"name": "Oct", "values": [412, 220]}], "unit": "USD"})
    card = box.picked[0]
    assert card["kind"] == "image" and card["image"].startswith(f"/api/files/{out['file_id']}?exp=")
    path, row = makers.find(mem, out["file_id"])
    assert path.read_bytes()[:4] == b"\x89PNG" and row["kind"] == "chart"
    q = dict(p.split("=") for p in card["image"].split("?")[1].split("&"))
    assert makers.link_ok(out["file_id"], q["exp"], q["sig"])
    assert not makers.link_ok(out["file_id"], q["exp"], "0" * 32)
    assert not makers.link_ok(out["file_id"], str(int(time.time()) - 1), makers.sign(out["file_id"], int(time.time()) - 1))


def test_word_document_and_deck_with_a_chart_inside(mem):
    chart, _ = run(mem, "make_chart", {"title": "Q", "type": "pie", "labels": ["a", "b"], "series": [{"values": [1, 3]}]})
    sections = [{"heading": "Summary", "paragraphs": ["All good."], "bullets": ["One", "Two"],
                 "table": {"columns": ["Item", "Cost"], "rows": [["Hotel", "200"]]}},
                {"heading": "Chart", "image_file_id": chart["file_id"], "notes": "Talk about it"}]
    doc, box = run(mem, "make_document", {"format": "docx", "title": "Trip plan", "sections": sections})
    deck, _ = run(mem, "make_document", {"format": "pptx", "title": "Trip plan", "subtitle": "Lima", "sections": sections})
    assert box.picked[0]["kind"] == "file" and "Word document" in box.picked[0]["subtitle"]
    from docx import Document
    from pptx import Presentation

    d = Document(str(makers.find(mem, doc["file_id"])[0]))
    assert "Summary" in [p.text for p in d.paragraphs] and len(d.tables) == 1 and len(d.inline_shapes) == 1
    p = Presentation(str(makers.find(mem, deck["file_id"])[0]))
    assert len(p.slides) == 3 and p.slides[2].notes_slide.notes_text_frame.text == "Talk about it"


def test_diagrams_are_checked_and_drawn(mem, monkeypatch):
    with pytest.raises(ToolError, match="images, links or file paths"):
        makers.render_diagram('digraph { a [image="/etc/shadow"]; }')
    with pytest.raises(ToolError):
        makers.render_diagram("print('hi')")
    seen = {}

    def fake_dot(args, input, capture_output, timeout):
        seen["dot"] = input.decode()
        return subprocess.CompletedProcess(args, 0, PNG, b"")
    if not shutil.which("dot"):
        monkeypatch.setattr(makers.subprocess, "run", fake_dot)
    out, box = run(mem, "make_diagram", {"title": "Flow", "dot": "digraph { rankdir=LR; Ask -> Think -> Answer; }"})
    assert box.picked[0]["kind"] == "image" and makers.find(mem, out["file_id"])[1]["kind"] == "diagram"
    if seen:
        assert "fillcolor" in seen["dot"]  # Buddy's styling was added


def test_images_refuse_personal_prompts_and_use_the_header_key(mem, monkeypatch):
    monkeypatch.setenv("BUDDY_GEMINI_KEY", "g-key")
    mem.remember("My daughter Priya Nair loves horses")
    with pytest.raises(PersonalData):
        asyncio.run(makers.run_maker_tool(mem, "make_image", {"prompt": "Priya Nair riding a horse"}))
    sent = {}

    def http(request):
        sent["headers"], sent["url"] = dict(request.headers), request.full_url
        return {"candidates": [{"content": {"parts": [{"inlineData": {"mimeType": "image/png",
                                                                      "data": base64.b64encode(PNG).decode()}}]}}]}
    assert makers.generate_image("a lighthouse at dawn, watercolour", http=http) == PNG
    assert sent["headers"]["X-goog-api-key"] == "g-key" and "key=" not in sent["url"]


def test_file_links_open_only_with_a_valid_signature(mem, monkeypatch):
    from fastapi.testclient import TestClient

    from buddy import server

    monkeypatch.setattr(server, "memory", mem)
    monkeypatch.setattr(server, "authorized", lambda t: t == "s3cret")
    file = makers.save(mem, PNG, "pic", "png", "image")
    url = makers.link(file["id"])
    with TestClient(server.app) as c:
        ok = c.get(url)
        assert ok.status_code == 200 and ok.headers["content-type"] == "image/png"
        assert c.get(f"/api/files/{file['id']}?exp=9999999999&sig=bad").status_code == 401
        assert c.get(f"/api/files/{file['id']}", headers={"Authorization": "Bearer s3cret"}).status_code == 200
        assert c.get("/api/files/../../etc/passwd").status_code in (401, 404)
