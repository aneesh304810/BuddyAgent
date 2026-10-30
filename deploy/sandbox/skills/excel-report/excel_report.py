#!/usr/bin/env python3
"""CSV/XLSX -> Excel report: a Summary sheet (totals per group + chart) and a Data sheet."""
import argparse

import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("data")
    p.add_argument("out")
    p.add_argument("--title", default="Report")
    p.add_argument("--group-by")
    p.add_argument("--sum")
    p.add_argument("--chart", choices=["bar", "pie", "none"], default="bar")
    a = p.parse_args()
    df = pd.read_excel(a.data) if a.data.endswith(".xlsx") else pd.read_csv(a.data)
    with pd.ExcelWriter(a.out, engine="openpyxl") as xl:
        if a.group_by and a.sum:
            summary = (df.groupby(a.group_by)[a.sum].sum().sort_values(ascending=False).reset_index())
            summary.loc[len(summary)] = ["Total", summary[a.sum].sum()]
            summary.to_excel(xl, sheet_name="Summary", index=False, startrow=2)
            ws = xl.sheets["Summary"]
            ws["A1"] = a.title
            ws["A1"].font = Font(size=16, bold=True)
            for cell in ws[3]:
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4C78DB")
            ws.cell(row=len(summary) + 3, column=1).font = Font(bold=True)
            ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 16
            if a.chart != "none" and len(summary) > 2:
                chart = PieChart() if a.chart == "pie" else BarChart()
                chart.title = a.title
                rows = len(summary) + 2  # without the Total row
                chart.add_data(Reference(ws, min_col=2, min_row=3, max_row=rows), titles_from_data=True)
                chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=rows))
                chart.height, chart.width = 9, 16
                ws.add_chart(chart, "D3")
        df.to_excel(xl, sheet_name="Data", index=False)
        data = xl.sheets["Data"]
        data.freeze_panes = "A2"
        data.auto_filter.ref = data.dimensions
        for i, col in enumerate(df.columns, 1):
            width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
            data.column_dimensions[get_column_letter(i)].width = min(width, 50)
    print(f"Wrote {a.out}")


if __name__ == "__main__":
    main()
